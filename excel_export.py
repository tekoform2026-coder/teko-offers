"""
excel_export.py
Количествена сметка (КСС) и чертежи в Excel за редактиране - по образеца reference/expected/12345676.xlsx:
  - лист „Кол. сметка“: всички артикули в реда на офертата, количествата са формули към листовете на елементите;
  - лист „Елемент N“: изглед отгоре и изглед отпред на всяка страна от слети клетки (1 клетка = 5 см),
    всеки панел и ъгъл с дебела външна граница; вдясно от чертежите - пълната КСС за елемента
    (за 1 бр. и общо = за 1 бр. x брой);
  - лист „Елементи Plaspanel“: всички панели и ъгли като готови блокове за копиране при редакция.
Разположението идва от teko_calculator.layout_element, геометрията в план - от drawing_generator.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from drawing_generator import PANEL_T, _face_strips, _plan_geometry
from teko_calculator import format_panel_code, is_compensator_row

CELL_CM = 5          # една клетка = 5 см
FIRST_COL = 3        # чертежите започват от колона C
COL_WIDTH = 2.6      # почти квадратни клетки
ROW_HEIGHT = 15

FILL_PANEL = "E6B9B8"
FILL_COMPENSATOR = "C8D1DC"
FILL_EX = "92D050"
FILL_IN = "00B050"
FILL_CONCRETE = "FAFAFA"
FILL_TITLE = "D9D9D9"
FILL_HEADER = "F2F2F2"
SIDE_COLOR = "1F4E79"

_black = Side(style="thin")
BORDER_KSS = Border(left=_black, right=_black, top=_black, bottom=_black)
_thick = Side(style="thick", color="000000")
HIDE_ZERO = "0;-0;;@"   # нулите в КСС не се показват
FILL_COUNT = "FFF2CC"

# Редът на артикулите в КСС (като в офертата) и кодовете от сметката, които влизат във всеки ред
KSS_ITEMS = [
    ("Панел TEKO  TK 150/60", ["TK 150/60", "TK 150/60 (полегнал)"]),
    ("Панел TEKO  TK 150/35", ["TK 150/35"]),
    ("Панел TEKO  TK 150/30", ["TK 150/30"]),
    ("Панел TEKO  TK 150/25", ["TK 150/25"]),
    ("Панел TEKO  TK 150/20", ["TK 150/20"]),
    ("Панел TEKO  TK 120/60", ["TK 120/60", "TK 120/60 (полегнал)"]),
    ("Панел TEKO  TK 120/35", ["TK 120/35"]),
    ("Панел TEKO  TK 120/30", ["TK 120/30"]),
    ("Панел TEKO  TK 120/25", ["TK 120/25"]),
    ("Панел TEKO  TK 120/20", ["TK 120/20"]),
    ("Панел TEKO  TK 60/35", ["TK 60/35"]),
    ("Панел TEKO  TK 60/30", ["TK 60/30"]),
    ("Панел TEKO  TK 60/25", ["TK 60/25"]),
    ("Панел TEKO  TK 60/20", ["TK 60/20"]),
    ("Панел TEKO  TK 60/15", ["TK 60/15"]),
    ("Панел TEKO  TK 60/10", ["TK 60/10"]),
    ("Ръкохватка H8", ["Ръкохватка H8"]),
    ("TEKO КОМПЕНСАТОР  TC150/15", ["TC 150/15"]),
    ("TEKO КОМПЕНСАТОР  TC150/10", ["TC 150/10"]),
    ("TEKO КОМПЕНСАТОР  TC150/5", ["TC 150/5"]),
    ("TEKO КОМПЕНСАТОР  TC120/15", ["TC 120/15"]),
    ("TEKO КОМПЕНСАТОР  TC120/10", ["TC 120/10"]),
    ("TEKO КОМПЕНСАТОР  TC120/5", ["TC 120/5"]),
    ("TEKO КОМПЕНСАТОР  TC60/15", ["TC 60/15"]),
    ("TEKO КОМПЕНСАТОР  TC60/10", ["TC 60/10"]),
    ("TEKO КОМПЕНСАТОР  TC60/5", ["TC 60/5"]),
    ("ВЪТРЕШЕН ЪГЪЛ  IN150", ["Вътрешен ъгъл IN150"]),
    ("ВЪТРЕШЕН ЪГЪЛ  IN120", ["Вътрешен ъгъл IN120"]),
    ("ВЪНШЕН ЪГЪЛ  EX240", ["Външен ъгъл EX240"]),
    ("ВЪНШЕН  СКОСЕН ЪГЪЛ EС240", []),
    ("Ръкохватка  H4", ["Ръкохватка H4"]),
    ("Подравняващ ригел   AW 100 поцинковани", ["Ригел AW 100"]),
    ("Подравняващ ригел   AW 200 поцинковани", ["Ригел AW 200"]),
    ("Подравняващ ригел   AW 300 поцинковани", ["Ригел AW 300"]),
    ("Вертикализатор PP", ["Вертикализатор PP"]),
    ("Панел TEKO  TK 40/25", ["TK 40/25"]),
    ("Панел TEKO  TK 30/25", ["TK 30/25"]),
    ("Панел TEKO  TK 20/20", ["TK 20/20"]),
    ("Панел TEKO  TK 20/15", ["TK 20/15"]),
    ("Панел TEKO  TK 15/15", ["TK 15/15"]),
    ("Ръкохватка  H6", ["Ръкохватка H6"]),
    ("Ръкохватка  H5", ["Ръкохватка H5"]),
    ("Шпилка TR 12", ["Шпилка TR 12"]),
    ("Гайка N12", ["Гайка N12"]),
]
_KSS_NAME_BY_CODE = {code: name for name, codes in KSS_ITEMS for code in codes}

TYPE_WORDS = {"column": "колона", "wall": "права стена", "l_wall": "L-стена",
              "u_wall": "П-стена", "t_wall": "T-стена"}


def kss_rows(bom):
    """
    Сметката {код: брой} по редовете на КСС: [(наименование, брой)] в реда на офертата.
    Кодове извън списъка (ако се появят) отиват най-отдолу със собственото си име.
    """
    totals = {}
    for code, qty in bom.items():
        name = _KSS_NAME_BY_CODE.get(code, code)
        totals[name] = totals.get(name, 0) + qty
    order = [name for name, _codes in KSS_ITEMS]
    extra = sorted(n for n in totals if n not in order)
    return [(n, totals[n]) for n in order + extra if totals.get(n)]


def _panel_label(width_cm, height_cm):
    """Надпис на панел в изгледа отпред: 150/60, 60/30, TC 150/15 (без „TK“)."""
    code = format_panel_code(int(width_cm), int(height_cm))
    return code.replace(" (полегнал)", "").replace("TK ", "")


def _cells(cm):
    return max(1, int(round(cm / CELL_CM)))


class _Sheet:
    """
    Рисуване с блокове от слети клетки; застъпващи се блокове не се сливат (Excel не ги приема).
    Блокът (панел, ъгъл) е с дебела външна граница, вътре - без линии.
    """

    def __init__(self, ws):
        self.ws = ws
        self.used = set()
        self.max_row = 1
        self.max_col = 1
        ws.sheet_format.defaultColWidth = COL_WIDTH
        ws.sheet_format.defaultRowHeight = ROW_HEIGHT
        ws.sheet_view.showGridLines = True

    def block(self, row, col, n_rows, n_cols, text=None, fill=None, size=13, rotation=0,
              bold=True, border=True, wrap=False, color=None):
        cells = [(r, c) for r in range(row, row + n_rows) for c in range(col, col + n_cols)]
        free = [rc for rc in cells if rc not in self.used]
        if not free:
            return
        self.used.update(free)
        fill_obj = PatternFill("solid", fgColor=fill) if fill else None
        last_row, last_col = row + n_rows - 1, col + n_cols - 1
        for r, c in free:
            cell = self.ws.cell(r, c)
            if fill_obj:
                cell.fill = fill_obj
            if border:
                cell.border = Border(left=_thick if c == col else None, right=_thick if c == last_col else None,
                                     top=_thick if r == row else None, bottom=_thick if r == last_row else None)
        if len(free) == len(cells) and len(cells) > 1:
            self.ws.merge_cells(start_row=row, start_column=col, end_row=last_row, end_column=last_col)
        r0, c0 = free[0]
        cell = self.ws.cell(r0, c0)
        if text is not None:
            cell.value = text
            cell.font = Font(bold=bold, size=size, color=color)
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       text_rotation=rotation, wrap_text=wrap)
        self.max_row = max(self.max_row, last_row)
        self.max_col = max(self.max_col, last_col)

    def cell(self, row, col, n_cols, value, bold=False, align="center", fill=None, number_format=None):
        """Клетка от таблица (слята по ширина) с тънка рамка; връща адреса ѝ."""
        self.ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + n_cols - 1)
        for c in range(col, col + n_cols):
            self.ws.cell(row, c).border = BORDER_KSS
            if fill:
                self.ws.cell(row, c).fill = PatternFill("solid", fgColor=fill)
            self.used.add((row, c))
        cell = self.ws.cell(row, col)
        cell.value = value
        cell.font = Font(size=10, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if number_format:
            cell.number_format = number_format
        self.max_row = max(self.max_row, row)
        self.max_col = max(self.max_col, col + n_cols - 1)
        return f"{get_column_letter(col)}{row}"

    def fill(self, row, col, color):
        if (row, col) in self.used:
            return
        self.ws.cell(row, col).fill = PatternFill("solid", fgColor=color)

    def text(self, row, col, value, size=8, bold=False, color="555555"):
        cell = self.ws.cell(row, col)
        cell.value = value
        cell.font = Font(size=size, bold=bold, color=color)
        self.max_row = max(self.max_row, row)
        self.max_col = max(self.max_col, col)


def _corner_block(sheet, row, col, n_rows, kind):
    sheet.block(row, col, n_rows, 2, kind, FILL_EX if kind == "EX" else FILL_IN, size=7, wrap=True)


# ==========================================
# ИЗГЛЕД ОТГОРЕ
# ==========================================
def _point_in_poly(x, y, poly):
    inside = False
    j = len(poly) - 1
    for i in range(len(poly)):
        xi, yi = poly[i]
        xj, yj = poly[j]
        if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi) + xi:
            inside = not inside
        j = i
    return inside


def _top_view(sheet, layout, top_row):
    """Изглед отгоре от долния ред панели на всяка страна; връща последния използван ред."""
    T = PANEL_T
    poly, place, corners = _plan_geometry(layout)
    strips = _face_strips(layout, place, poly)

    rects = []   # (x, y, w, h, текст, цвят, завъртане, размер на шрифта)
    for face in layout["faces"]:
        x0, y0, axis, *_ = place[face["name"]]
        s = 0
        for w in face["rows"][0]:
            if axis == 'x':
                rects.append((x0 + s, y0, w, T, str(int(w)), FILL_PANEL, 0, 13))
            else:
                rects.append((x0, y0 + s, T, w, str(int(w)), FILL_PANEL, 90, 13))
            s += w
    for kind, x, y in corners:
        rects.append((x, y, T, T, kind, FILL_EX if kind == "EX" else FILL_IN, 0, 7))

    # Букви на страните - една клетка навън от панела, в средата на страната
    letters = []
    for name, (axis, a0, a1, c_lo, c_hi, normal) in strips.items():
        if name.startswith("Край"):
            continue
        cross = (c_hi + CELL_CM) if normal > 0 else (c_lo - 2 * CELL_CM)
        along = (a0 + a1) / 2 - CELL_CM / 2
        x, y = (along, cross) if axis == 'x' else (cross, along)
        letters.append((x, y, name.split()[0]))

    xs = [r[0] for r in rects] + [r[0] + r[2] for r in rects] + [p[0] for p in poly] + \
         [v for x, _, _ in letters for v in (x, x + CELL_CM)]
    ys = [r[1] for r in rects] + [r[1] + r[3] for r in rects] + [p[1] for p in poly] + \
         [v for _, y, _ in letters for v in (y, y + CELL_CM)]
    x_min, y_max = min(xs), max(ys)

    def col_of(x):
        return FIRST_COL + int(round((x - x_min) / CELL_CM))

    def row_of(y_top):
        return top_row + int(round((y_max - y_top) / CELL_CM))

    for x, y, w, h, text, color, rot, size in rects:
        sheet.block(row_of(y + h), col_of(x), _cells(h), _cells(w), text, color, size=size,
                    rotation=rot, wrap=(size == 7))

    # бетон: клетките със среда вътре в контура
    pxs, pys = [p[0] for p in poly], [p[1] for p in poly]
    for r in range(row_of(max(pys)), row_of(min(pys))):
        for c in range(col_of(min(pxs)), col_of(max(pxs))):
            cx = x_min + (c - FIRST_COL + 0.5) * CELL_CM
            cy = y_max - (r - top_row + 0.5) * CELL_CM
            if _point_in_poly(cx, cy, poly):
                sheet.fill(r, c, FILL_CONCRETE)

    for x, y, label in letters:
        r, c = row_of(y + CELL_CM), col_of(x)
        sheet.block(r, c, 1, 1, label, None, size=9, border=False, color=SIDE_COLOR)

    return max(row_of(min(ys)) - 1, top_row)


# ==========================================
# ИЗГЛЕД ОТПРЕД
# ==========================================
def _face_groups(layout):
    """Страните за изгледите отпред: външните, после вътрешните, после челните; еднаквите - заедно."""
    def kind(name):
        if name.startswith("Край"):
            return 2
        return 1 if name.split()[0].endswith("1") else 0

    groups = []
    for face in sorted(layout["faces"], key=lambda f: kind(f["name"])):
        for g in groups:
            ref = g[0]
            if (ref["rows"], ref["start"], ref["end"]) == (face["rows"], face["start"], face["end"]):
                g.append(face)
                break
        else:
            groups.append([face])
    # еднаквите страни се показват на мястото на първата от тях
    return sorted(groups, key=lambda g: kind(g[0]["name"]))


def _front_view(sheet, layout, face, top_row):
    """Изглед отпред на една страна: редовете отгоре надолу, ъглите в краищата по цялата височина."""
    levels = layout["height_levels"]
    total_rows = sum(_cells(h) for h in levels)
    col = FIRST_COL
    if face["start"]:
        _corner_block(sheet, top_row, col, total_rows, face["start"])
        col += 2
    row = top_row
    end_col = col
    for h, widths in reversed(list(zip(levels, face["rows"]))):
        c = col
        for w in widths:
            if is_compensator_row(h):
                sheet.block(row, c, _cells(h), _cells(w), f"TC {int(h)} см", FILL_COMPENSATOR, size=9)
            else:
                sheet.block(row, c, _cells(h), _cells(w), _panel_label(w, h), FILL_PANEL, size=14,
                            rotation=90 if h > w else 0)
            c += _cells(w)
        end_col = max(end_col, c)
        row += _cells(h)
    if face["end"]:
        _corner_block(sheet, top_row, end_col, total_rows, face["end"])
    return top_row + total_rows - 1


# ==========================================
# ЛИСТ НА ЕЛЕМЕНТ
# ==========================================
def _element_title(index, element):
    layout = element["layout"]
    d = layout["dims"]
    typ = layout["type"]
    parts = [str(element.get("name") or "").strip(), TYPE_WORDS.get(typ, typ)]
    if typ == "column":
        parts.append(f"{d['a']:g}x{d['b']:g}")
    elif typ == "wall":
        parts.append(f"дължина {d['length']:g}")
    elif typ == "u_wall":
        parts += [f"A {d['l1']:g}", f"B {d['l2']:g}", f"C {d['l3']:g}"]
    else:
        parts += [f"A {d['l1']:g}", f"B {d['l2']:g}"]
    parts.append(f"височина {layout['height_cm']:g}")
    if typ != "column":
        parts.append(f"дебелина {layout['thickness_cm']:g}")
    parts.append(f"бр. {element.get('count', 1)}")
    return f"Елемент {index}  —  " + " · ".join(p for p in parts if p)


def _element_sheet(wb, index, element):
    """Лист „Елемент N“; връща {наименование в КСС: адрес на клетката „Кол. общо“}."""
    ws = wb.create_sheet(f"Елемент {index}")
    sheet = _Sheet(ws)
    layout = element["layout"]
    count = max(1, int(element.get("count", 1)))

    # Изглед отгоре, под него изгледите отпред
    sheet.text(4, FIRST_COL, "Изглед отгоре")
    last = _top_view(sheet, layout, 5)
    for group in _face_groups(layout):
        label_row = last + 3
        names = " / ".join(f["name"] for f in group)
        sheet.text(label_row, FIRST_COL, f"Изглед — {names}")
        last = _front_view(sheet, layout, group[0], label_row + 1)

    # КСС за елемента вдясно от чертежите: всички артикули като в „Кол. сметка“
    name_col = sheet.max_col + 3
    unit_col = name_col + 16
    one_col = unit_col + 3
    all_col = one_col + 4
    sheet.cell(4, name_col, all_col - name_col, "Брой елементи:", bold=True, align="right")
    sheet.cell(4, all_col, 4, count, bold=True, fill=FILL_COUNT)
    count_ref = f"${get_column_letter(all_col)}$4"
    ws.row_dimensions[5].height = 30
    for col, n, title in ((name_col, 16, "Елемент"), (unit_col, 3, "Мярка"),
                          (one_col, 4, "Кол. за 1 бр."), (all_col, 4, "Кол. общо")):
        sheet.cell(5, col, n, title, bold=True, fill=FILL_HEADER)

    totals = dict(kss_rows(element["bom"]))
    order = [name for name, _codes in KSS_ITEMS]
    refs = {}
    for i, name in enumerate(order + sorted(n for n in totals if n not in order)):
        r = 6 + i
        one = totals.get(name, 0) / count
        sheet.cell(r, name_col, 16, name, align="left")
        sheet.cell(r, unit_col, 3, "бр")
        one_ref = sheet.cell(r, one_col, 4, (int(one) if one == int(one) else one) or None, number_format=HIDE_ZERO)
        refs[name] = sheet.cell(r, all_col, 4, f'=IF({one_ref}="","",{one_ref}*{count_ref})', bold=True,
                                number_format=HIDE_ZERO)

    ws.merge_cells(start_row=2, start_column=FIRST_COL, end_row=2, end_column=max(sheet.max_col, FIRST_COL + 20))
    title = ws.cell(2, FIRST_COL)
    title.value = _element_title(index, element)
    title.font = Font(bold=True, size=11)
    title.alignment = Alignment(vertical="center")
    title.fill = PatternFill("solid", fgColor=FILL_TITLE)
    ws.row_dimensions[2].height = 18
    return refs


# ==========================================
# КОЛИЧЕСТВЕНА СМЕТКА
# ==========================================
def _kss_sheet(ws, names, quantity):
    """Лист КСС като „Кол. сметка“ в образеца: наименование, мярка, количество (формула или число)."""
    ws.column_dimensions["A"].width = 19
    ws.column_dimensions["B"].width = 41.43
    ws.column_dimensions["C"].width = 7.43
    ws.column_dimensions["D"].width = 9.86
    ws.row_dimensions[4].height = 23.25
    ws.row_dimensions[5].height = 45
    ws.row_dimensions[6].height = 15.75
    for col, value in (("B", "Елемент"), ("C", "Мярка"), ("D", "Кол.")):
        ws[f"{col}5"] = value
    ws["D6"] = "[бр]"
    for r in (5, 6):
        for col in "BCD":
            cell = ws[f"{col}{r}"]
            cell.font = Font(size=12)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER_KSS

    for i, name in enumerate(names):
        r = 7 + i
        ws.row_dimensions[r].height = 18
        ws[f"B{r}"] = name
        ws[f"C{r}"] = "бр"
        qty = quantity(name)
        if qty:
            ws[f"D{r}"] = qty
            ws[f"D{r}"].number_format = HIDE_ZERO
        for col in "BCD":
            cell = ws[f"{col}{r}"]
            cell.font = Font(size=11)
            cell.border = BORDER_KSS
            if col != "B":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")


# ==========================================
# ПАЛИТРА С ЕЛЕМЕНТИ
# ==========================================
PALETTE_STANDING = [(150, w) for w in (60, 35, 30, 25, 20, 15, 10, 5)] + \
                   [(120, w) for w in (60, 35, 30, 25, 20, 15, 10, 5)] + \
                   [(60, w) for w in (35, 30, 25, 20, 15, 10, 5)] + \
                   [(40, 25), (30, 25), (20, 20), (20, 15), (15, 15)]


def _palette_sheet(wb):
    """Всички панели и ъгли: отгоре - изглед отгоре, отдолу - изглед отпред; после легналите."""
    ws = wb.create_sheet("Елементи Plaspanel")
    sheet = _Sheet(ws)
    sheet.text(4, FIRST_COL, "Елементи (1 клетка = 5 см) - копирайте блок и го поставете в чертежа",
               size=11, bold=True, color="000000")
    sheet.text(6, FIRST_COL, "Изправени панели", size=9, bold=True)

    col = FIRST_COL
    for h, w in PALETTE_STANDING:
        sheet.block(7, col, 2, _cells(w), str(w), FILL_PANEL, size=9)
        sheet.block(12, col, _cells(h), _cells(w), _panel_label(w, h), FILL_PANEL, size=11,
                    rotation=90 if h > w else 0)
        col += _cells(w) + 2
    for label, kind, height in (("IN 150", "IN", 150), ("IN 120", "IN", 120), ("EX 240", "EX", 240)):
        sheet.block(7, col, 2, 2, kind, FILL_EX if kind == "EX" else FILL_IN, size=7, wrap=True)
        sheet.block(12, col, _cells(height), 2, label, FILL_EX if kind == "EX" else FILL_IN, size=9, rotation=90)
        col += 4

    row = 12 + _cells(240) + 4
    sheet.text(row - 1, FIRST_COL, "Легнали панели (изглед отгоре вляво, изглед отпред вдясно)", size=9, bold=True)
    for h, w in [p for p in PALETTE_STANDING if p[0] in (150, 120)]:
        sheet.block(row, FIRST_COL, _cells(w), 2, str(w), FILL_PANEL, size=9, rotation=90)
        sheet.block(row, FIRST_COL + 4, _cells(w), _cells(h), _panel_label(w, h), FILL_PANEL, size=11)
        row += _cells(w) + 2


# ==========================================
# ГЛАВНА ФУНКЦИЯ
# ==========================================
def generate_excel(elements):
    """
    elements: [{"name", "count", "layout", "bom"}] - bom е сметката на елемента, умножена по броя.
    Връща BytesIO с .xlsx.
    """
    wb = Workbook()
    kss = wb.active
    kss.title = "Кол. сметка"
    element_refs = []
    extra = []
    known = {name for name, _codes in KSS_ITEMS}
    for i, element in enumerate(elements, 1):
        refs = _element_sheet(wb, i, element)
        element_refs.append((f"Елемент {i}", refs))
        extra += [n for n in refs if n not in known and n not in extra]
    def formula(name):
        sources = [f"'{sheet}'!{refs[name]}" for sheet, refs in element_refs if name in refs]
        return f"=SUM({','.join(sources)})" if sources else None

    _kss_sheet(kss, [name for name, _codes in KSS_ITEMS] + extra, formula)
    _palette_sheet(wb)
    wb.calculation.fullCalcOnLoad = True

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
